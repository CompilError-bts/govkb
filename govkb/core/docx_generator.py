from __future__ import annotations

import re
from pathlib import Path

from govkb.models.article import Article
from govkb.models.site_info import SiteInfo


class DocxGenerator:
    """Generate one Word document per article."""

    TEMPLATE_CONFIG = {
        "title": {"font": "SimHei", "size": 16, "bold": True, "align": "center"},
        "meta": {"font": "SimSun", "size": 10, "color": "808080", "align": "center"},
        "body": {"font": "SimSun", "size": 12, "indent": 24, "line_spacing": 1.5},
        "footer": {"font": "SimSun", "size": 9, "color": "808080", "align": "center"},
    }

    def __init__(self, site: SiteInfo, *, filename_pattern: str = "{index:03d}_{date}_{title}.docx"):
        self.site = site
        self.filename_pattern = filename_pattern
        self.output_dir: Path | None = None
        self._current_index = 0

    def generate(self, articles: list[Article], output_dir: str) -> int:
        """Generate docx files in output_dir and return the success count."""

        self._ensure_docx_available()
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        success = 0
        for index, article in enumerate(articles, start=1):
            self._current_index = index
            try:
                self.generate_one(article)
                success += 1
            except Exception:
                continue
        return success

    def generate_one(self, article: Article) -> str:
        """Generate a docx file for one article and return its path."""

        self._ensure_docx_available()
        if self.output_dir is None:
            self.output_dir = Path.cwd() / "output"
            self.output_dir.mkdir(parents=True, exist_ok=True)

        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        doc = Document()
        self._add_title(doc, article.title, WD_ALIGN_PARAGRAPH, Pt, qn)
        self._add_meta(doc, article, WD_ALIGN_PARAGRAPH, Pt, RGBColor, qn)
        self._add_divider(doc, Pt, RGBColor, qn)
        self._add_body(doc, article.content, Pt, qn)
        self._add_footer(doc, Pt, RGBColor, qn)

        file_path = self._next_file_path(article)
        doc.save(file_path)
        return str(file_path)

    def _add_title(self, doc, title: str, align, pt, qn) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = align.CENTER
        run = paragraph.add_run(title)
        self._set_font(run, "SimHei", 16, pt, qn)
        run.bold = True

    def _add_meta(self, doc, article: Article, align, pt, rgb_color, qn) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = align.CENTER
        run = paragraph.add_run(f"发布日期：{article.date}    来源：{self.site.portal_name}")
        self._set_font(run, "SimSun", 10, pt, qn)
        run.font.color.rgb = rgb_color(0x80, 0x80, 0x80)

        url_paragraph = doc.add_paragraph()
        url_paragraph.alignment = align.CENTER
        url_run = url_paragraph.add_run(article.url)
        self._set_font(url_run, "SimSun", 9, pt, qn)
        url_run.font.color.rgb = rgb_color(0x80, 0x80, 0x80)

    def _add_divider(self, doc, pt, rgb_color, qn) -> None:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("-" * 50)
        self._set_font(run, "SimSun", 9, pt, qn)
        run.font.color.rgb = rgb_color(0xB0, 0xB0, 0xB0)

    def _add_body(self, doc, content: str, pt, qn) -> None:
        paragraphs = [line.strip() for line in re.split(r"\n+", content or "") if line.strip()]
        if not paragraphs:
            paragraphs = ["[正文为空]"]

        for text in paragraphs:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = pt(24)
            paragraph.paragraph_format.line_spacing = 1.5
            run = paragraph.add_run(text)
            self._set_font(run, "SimSun", 12, pt, qn)

    def _add_footer(self, doc, pt, rgb_color, qn) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = 1
        run = paragraph.add_run(f"来源：{self.site.portal_name}（{self.site.domain}）")
        self._set_font(run, "SimSun", 9, pt, qn)
        run.font.color.rgb = rgb_color(0x80, 0x80, 0x80)

    def _set_font(self, run, font_name: str, size: int, pt, qn) -> None:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        run.font.size = pt(size)

    def _next_file_path(self, article: Article) -> Path:
        safe_title = self._safe_filename(article.title, max_length=60)
        filename = self.filename_pattern.format(
            index=max(1, self._current_index),
            date=article.date,
            title=safe_title,
        )
        path = self.output_dir / filename  # type: ignore[operator]
        if not path.exists():
            return path

        stem = path.stem
        suffix = path.suffix
        counter = 2
        while True:
            candidate = path.with_name(f"{stem}_{counter}{suffix}")
            if not candidate.exists():
                return candidate
            counter += 1

    def _safe_filename(self, value: str, max_length: int = 80) -> str:
        value = re.sub(r'[\\/:*?"<>|]', "_", value)
        value = re.sub(r"\s+", " ", value).strip(" .")
        if len(value) > max_length:
            value = value[:max_length].rstrip(" .")
        return value or "untitled"

    def _ensure_docx_available(self) -> None:
        try:
            import docx  # noqa: F401
        except ModuleNotFoundError as exc:
            raise RuntimeError("Missing dependency: python-docx. Install it with `pip install python-docx`.") from exc
